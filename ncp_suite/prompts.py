"""Load the prompt sheet (data/mcp_prompts.xlsx, sheet 'prompts').

CHANGED: optional `Timeout` column (seconds per prompt). Blank = the keyword rule in
chat/policy.py, as before.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class PromptRow:
    id: str
    prompt: str
    check: str
    param: float | None = None        # e.g. the 80 in "CPU above 80%"
    tol: float | None = None          # overrides the default tolerance / min coverage
    applies_to: frozenset = frozenset()   # connector keys; empty = all
    known_issue: str = ""
    timeout: int | None = None        # seconds to wait for NCP's answer; None = keyword rule

    def applies(self, connector: str) -> bool:
        return not self.applies_to or connector in self.applies_to


def _f(v) -> float | None:
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def load_prompts(path: Path) -> list[PromptRow]:
    ws = load_workbook(path, read_only=True, data_only=True)["prompts"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows)]
    out = []
    for values in rows:
        r = dict(zip(header, values))
        if not r.get("id") or not r.get("prompt"):
            continue
        applies = {x.strip().lower() for x in str(r.get("applies_to") or "").split(",") if x.strip()}
        timeout = _f(r.get("timeout"))
        out.append(PromptRow(
            id=str(r["id"]).strip(), prompt=str(r["prompt"]).strip(), check=str(r.get("check") or "").strip(),
            param=_f(r.get("param")), tol=_f(r.get("tolerance")), applies_to=frozenset(applies - {"all"}),
            known_issue=str(r.get("known_issue") or "").strip(),
            timeout=int(timeout) if timeout else None,
        ))
    return out
